"""Bank statements (xlsx) with per-row ground truth: category and recurring group."""
from __future__ import annotations

import calendar
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .company import BANKS, Company, account, counterparty
from .translit import to_cyrillic

CATEGORIES = ["sales_receipt", "supplier_payment", "salary", "payroll_tax", "tax", "loan_repayment",
              "loan_disbursement", "rent_utilities", "own_transfer", "cash_withdrawal", "bank_fee", "other"]
MONTH_RU = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь",
            "ноябрь", "декабрь"]
MONTH_RU_PREP = ["январе", "феврале", "марте", "апреле", "мае", "июне", "июле", "августе", "сентябре",
                 "октябре", "ноябре", "декабре"]
MONTH_UZ = ["yanvar", "fevral", "mart", "aprel", "may", "iyun", "iyul", "avgust", "sentyabr", "oktyabr",
            "noyabr", "dekabr"]
TAX_OFFICE = ("Davlat soliq qo'mitasi (budjet hisobvarag'i)", "201122905")
FUND = ("Budjetdan tashqari Pensiya jamg'armasi", "201122919")


@dataclass
class Party:
    name: str
    stir: str | None
    account: str
    mfo: str


@dataclass
class Txn:
    d: date
    amount: int                 # UZS, + inflow, - outflow
    party: Party
    purpose: str
    category: str
    group: str | None = None    # candidate recurring group key
    doc_no: int = 0
    recurring_group: str | None = None
    meta: dict = field(default_factory=dict)


def party_from(rng, name: str, stir: str | None) -> Party:
    _bank, mfo = BANKS[int(rng.integers(0, len(BANKS)))]
    return Party(name, stir, account(rng, "20208"), mfo)


def business_day(rng, y: int, m: int, day: int | None = None, spread: int = 0) -> date:
    last = calendar.monthrange(y, m)[1]
    if day is None:
        day = int(rng.integers(1, last + 1))
    day = int(np.clip(day + (int(rng.integers(-spread, spread + 1)) if spread else 0), 1, last))
    d = date(y, m, day)
    while d.weekday() >= 5:  # move weekend payments to Monday (or back to Friday at month end)
        nd = d.toordinal() + 1
        if date.fromordinal(nd).month != m:
            d = date.fromordinal(d.toordinal() - 1)
            while d.weekday() >= 5:
                d = date.fromordinal(d.toordinal() - 1)
            break
        d = date.fromordinal(nd)
    return d


class Grammar:
    """Purpose-text grammar. `lang` is the statement's dominant language (ru or uz)."""

    def __init__(self, rng: np.random.Generator, lang: str = "ru", typo_rate: float = 0.03):
        self.rng, self.lang, self.typo = rng, lang, typo_rate

    def _pick(self, ru: list[str], uz: list[str]) -> str:
        pool = ru if (self.lang == "ru") ^ (self.rng.random() < 0.2) else uz
        return str(self.rng.choice(pool))

    def _noise(self, s: str) -> str:
        r = self.rng
        if r.random() < 0.25:
            s = (s.replace("по договору", "по дог.").replace("счет-фактуре", "с/ф")
                 .replace("shartnoma", "shartn.").replace("hisob-faktura", "h/f"))
        if r.random() < self.typo and len(s) > 12:
            i = int(r.integers(3, len(s) - 3))
            s = s[:i] + s[i + 1:]
        if r.random() < 0.08:
            s = s.upper()
        return s

    def render(self, kind: str, **k) -> str:
        r = self.rng
        n = k.get("n", int(r.integers(10, 9999)))
        dd = k.get("d", date(2026, 1, 1)).strftime("%d.%m.%Y")
        c = k.get("contract", f"{int(r.integers(1, 300))}/{str(k.get('d', date(2026, 1, 1)).year)[2:]}")
        m = k.get("month", 1) - 1
        y = k.get("year", 2026)
        if kind == "sales_receipt":
            s = self._pick([
                f"Оплата за товар по счет-фактуре №{n} от {dd}, в т.ч. НДС 12%",
                f"Оплата за продукцию по договору №{c}, НДС 12%",
                f"Оплата по счет-фактуре №{n} за поставленный товар",
                f"Предоплата за товар по договору №{c} от {dd}",
            ], [
                f"Tovar uchun to'lov, hisob-faktura №{n} {dd}, QQS bilan",
                f"Mahsulot uchun to'lov, shartnoma №{c}, QQS 12%",
                f"Yetkazilgan tovar uchun to'lov h/f №{n}",
                f"Shartnoma №{c} bo'yicha oldindan to'lov",
            ])
        elif kind == "supplier_payment":
            s = self._pick([
                f"Оплата за сырье по договору №{c} от {dd}, НДС 12%",
                f"Оплата за товары согласно счет-фактуре №{n}",
                f"Предоплата за материалы по договору №{c}",
                f"Оплата за транспортные услуги по договору №{c}",
                f"Оплата за упаковочные материалы, с/ф №{n}",
            ], [
                f"Xom ashyo uchun to'lov, shartnoma №{c}, QQS bilan",
                f"Tovarlar uchun to'lov h/f №{n}",
                f"Materiallar uchun oldindan to'lov, shartnoma №{c}",
                f"Transport xizmatlari uchun to'lov, shartnoma №{c}",
            ])
        elif kind == "salary":
            s = self._pick([
                f"Заработная плата за {MONTH_RU[m]} {y} г. (перечисление на пластиковые карты)",
                f"Перечисление заработной платы за {MONTH_RU[m]} {y} г. согласно реестру №{n}",
            ], [
                f"Ish haqi to'lovi {MONTH_UZ[m]} {y} (plastik kartalarga)",
                f"{MONTH_UZ[m].capitalize()} oyi uchun ish haqi, reestr №{n}",
            ])
        elif kind == "social_tax":
            s = self._pick([f"Социальный налог за {MONTH_RU[m]} {y} г.", f"Уплата социального налога за {MONTH_RU[m]}"],
                           [f"Ijtimoiy soliq {MONTH_UZ[m]} {y}", f"{MONTH_UZ[m].capitalize()} oyi uchun ijtimoiy soliq"])
        elif kind == "pit":
            s = self._pick([f"Налог на доходы физических лиц за {MONTH_RU[m]} {y} г.", f"НДФЛ за {MONTH_RU[m]} {y}"],
                           [f"JShDS {MONTH_UZ[m]} oyi uchun", f"Jismoniy shaxslar daromad solig'i {MONTH_UZ[m]} {y}"])
        elif kind == "vat":
            s = self._pick([f"НДС за {MONTH_RU[m]} {y} г.", f"Уплата налога на добавленную стоимость за {MONTH_RU[m]}"],
                           [f"QQS {MONTH_UZ[m]} {y} uchun", f"Qo'shilgan qiymat solig'i {MONTH_UZ[m]} oyi"])
        elif kind == "profit_tax":
            s = self._pick(["Налог на прибыль, авансовый платеж", f"Авансовый платеж по налогу на прибыль за {MONTH_RU[m]}"],
                           ["Foyda solig'i bo'yicha bo'nak to'lov", f"Foyda solig'i bo'nak to'lovi {MONTH_UZ[m]}"])
        elif kind == "property_tax":
            s = self._pick(["Налог на имущество юридических лиц", "Земельный налог"],
                           ["Yuridik shaxslar mol-mulk solig'i", "Yer solig'i"])
        elif kind == "loan_repayment":
            cd = k.get("contract_date", "")
            s = self._pick([
                f"Погашение основного долга и процентов по кредитному договору №{c} от {cd}",
                f"Погашение кредита по договору № {c} от {cd}",
                f"Погашение задолженности по кредиту, дог. №{c}",
            ], [
                f"Kredit bo'yicha asosiy qarz va foizlarni qaytarish, shartnoma №{c}",
                f"Kreditni qaytarish, shartnoma №{c} {cd}",
            ])
        elif kind == "loan_interest":
            s = self._pick([f"Погашение процентов по кредиту, дог. №{c}"], [f"Kredit foizlarini to'lash, shartnoma №{c}"])
        elif kind == "loan_disbursement":
            s = self._pick([f"Выдача кредита по кредитному договору №{c}"], [f"Kredit mablag'lari ajratildi, shartnoma №{c}"])
        elif kind == "rent":
            s = self._pick([f"Арендная плата за {MONTH_RU[m]} {y} г. по договору №{c}", f"Аренда помещения за {MONTH_RU[m]}, дог. №{c}"],
                           [f"Ijara haqi {MONTH_UZ[m]} oyi uchun, shartnoma №{c}", f"Bino ijarasi uchun to'lov {MONTH_UZ[m]}"])
        elif kind == "electricity":
            s = self._pick([f"Оплата за электроэнергию за {MONTH_RU[m]} {y} г.", "Оплата за электроэнергию согласно счету"],
                           [f"Elektr energiyasi uchun to'lov {MONTH_UZ[m]}", "Elektr energiyasi uchun to'lov"])
        elif kind == "gas":
            s = self._pick(["Оплата за природный газ", f"Оплата за газ за {MONTH_RU[m]}"], ["Tabiiy gaz uchun to'lov"])
        elif kind == "water":
            s = self._pick(["Оплата за холодную воду и канализацию", "Оплата за водоснабжение"], ["Ichimlik suvi uchun to'lov"])
        elif kind == "own_transfer":
            s = self._pick(["Перевод собственных средств на другой счет", "Перевод собственных средств"],
                           ["O'z mablag'larini boshqa hisob raqamiga o'tkazish"])
        elif kind == "own_transfer_in":
            s = self._pick(["Поступление собственных средств с другого счета"], ["O'z mablag'lari boshqa hisobdan"])
        elif kind == "cash_withdrawal":
            s = self._pick([f"Получение наличных по чеку №{n} на хоз. нужды", "Снятие наличных на хозяйственные нужды"],
                           ["Xo'jalik ehtiyojlari uchun naqd pul olish", f"Chek №{n} bo'yicha naqd pul"])
        elif kind == "bank_fee":
            s = self._pick([f"Комиссия банка за расчетно-кассовое обслуживание за {MONTH_RU[m]}", "Плата за SMS-информирование",
                            "Комиссия за перевод"], ["Bank xizmati uchun komissiya", "Hisob-kassa xizmati uchun haq"])
        elif kind == "refund":
            s = self._pick(["Возврат ошибочно перечисленных средств", "Возврат предоплаты по договору"],
                           ["Xato o'tkazilgan mablag'ni qaytarish"])
        elif kind == "vague":
            s = self._pick([f"Оплата по договору №{c}", "Перевод средств", f"Оплата согласно письму №{n}", "Оплата по счету",
                            f"Платеж №{n}"], [f"Shartnoma №{c} bo'yicha to'lov", "Mablag' o'tkazish", f"Xat №{n} asosida to'lov"])
        elif kind == "other":
            s = self._pick(["Благотворительная помощь", "Членский взнос Торгово-промышленной палаты", f"Штраф по договору №{c}",
                            "Оплата за обучение сотрудников"], ["Homiylik yordami", "Savdo-sanoat palatasiga a'zolik badali"])
        else:
            raise ValueError(kind)
        s = self._noise(s)
        if self.lang == "uz_cyrl" and self.rng.random() < 0.5:
            s = to_cyrillic(s)
        return s


def assign_recurring(txns: list[Txn]) -> None:
    """Label recurring groups with the plan's definition: >=3 occurrences, day-of-month std <= 3, amount CV <= 0.15."""
    groups: dict[str, list[Txn]] = {}
    for t in txns:
        if t.group:
            groups.setdefault(t.group, []).append(t)
    for key, ts in groups.items():
        if len(ts) < 3:
            continue
        days = np.array([t.d.day for t in ts], dtype=float)
        amts = np.array([abs(t.amount) for t in ts], dtype=float)
        if days.std() <= 3 and amts.std() / max(amts.mean(), 1) <= 0.15:
            for t in ts:
                t.recurring_group = key


def write_statement(path: Path, company: Company, txns: list[Txn], opening: int, start: date, end: date,
                    layout: str = "A", lang: str = "ru") -> dict:
    """Write xlsx; returns truth dict with row numbers (1-based Excel rows)."""
    txns = sorted(txns, key=lambda t: (t.d, -t.amount))
    for i, t in enumerate(txns):
        t.doc_no = t.doc_no or (100 + i)
    wb = Workbook()
    ws = wb.active
    ws.title = "Выписка" if layout == "A" else "Ko'chirma"
    bold = Font(bold=True)
    money = "#,##0.00"
    rows_truth = []
    bal = opening
    month_end: dict[str, int] = {}
    if layout == "A":
        ws.append([company.bank]); ws["A1"].font = bold
        ws.append([f"ВЫПИСКА ПО РАСЧЕТНОМУ СЧЕТУ № {company.account}"]); ws["A2"].font = bold
        ws.append([f"Клиент: {company.name}", None, None, None, f"ИНН: {company.stir}"])
        ws.append([f"Период: с {start:%d.%m.%Y} по {end:%d.%m.%Y}"])
        ws.append([f"Входящий остаток на {start:%d.%m.%Y}:", None, opening]); ws["C5"].number_format = money
        ws.append([])
        ws.append(["Дата", "№ док.", "Дебет", "Кредит", "Наименование контрагента", "ИНН контрагента",
                   "Счет контрагента", "МФО банка", "Назначение платежа", "Остаток"])
        for cell in ws[7]:
            cell.font = bold
        header_rows, first = [7], 8
        for t in txns:
            bal += t.amount
            ws.append([t.d, t.doc_no, -t.amount if t.amount < 0 else None, t.amount if t.amount > 0 else None,
                       t.party.name, t.party.stir, t.party.account, t.party.mfo, t.purpose, bal])
            r = ws.max_row
            ws.cell(r, 1).number_format = "DD.MM.YYYY"
            for col in (3, 4, 10):
                ws.cell(r, col).number_format = money
            rows_truth.append((r, t))
            month_end[f"{t.d:%Y-%m}"] = bal
        deb = sum(-t.amount for t in txns if t.amount < 0)
        cre = sum(t.amount for t in txns if t.amount > 0)
        ws.append(["Итого обороты:", None, deb, cre])
        ws.append([f"Исходящий остаток на {end:%d.%m.%Y}:", None, bal])
        for r in (ws.max_row - 1, ws.max_row):
            for col in (3, 4):
                ws.cell(r, col).number_format = money
                ws.cell(r, col).font = bold
        widths = [12, 9, 16, 16, 38, 13, 24, 9, 70, 18]
    else:
        L = (lambda s: to_cyrillic(s)) if lang == "uz_cyrl" else (lambda s: s)
        ws.append([L(company.bank)]); ws["A1"].font = bold
        ws.append([L("HISOB RAQAMIDAN KO'CHIRMA")]); ws["A2"].font = bold
        ws.append([L(f"Hisob raqami: {company.account}")])
        ws.append([L(f"Mijoz: {company.name}, STIR {company.stir}")])
        ws.append([L(f"Davr: {start:%d.%m.%Y} - {end:%d.%m.%Y}")])
        ws.append([L("Davr boshidagi qoldiq:"), None, opening]); ws["C6"].number_format = money
        ws.append([])
        ws.append([L("Sana"), L("Hujjat №"), L("Aylanma"), None, L("Kontragent"), None, None, None, L("To'lov maqsadi")])
        ws.append([None, None, L("Debet"), L("Kredit"), L("Nomi"), L("STIR"), L("Hisob raqami"), L("MFO"), None])
        ws.merge_cells("C8:D8")
        ws.merge_cells("E8:H8")
        for rr in (8, 9):
            for cell in ws[rr]:
                cell.font = bold
                cell.alignment = Alignment(horizontal="center")
        header_rows, first = [8, 9], 10
        for t in txns:
            bal += t.amount
            ws.append([t.d.strftime("%d.%m.%Y"), str(t.doc_no), -t.amount if t.amount < 0 else None,
                       t.amount if t.amount > 0 else None, t.party.name, t.party.stir, t.party.account,
                       t.party.mfo, t.purpose])
            r = ws.max_row
            for col in (3, 4):
                ws.cell(r, col).number_format = money
            rows_truth.append((r, t))
            month_end[f"{t.d:%Y-%m}"] = bal
        deb = sum(-t.amount for t in txns if t.amount < 0)
        cre = sum(t.amount for t in txns if t.amount > 0)
        ws.append([L("Jami aylanma:"), None, deb, cre])
        ws.append([L("Davr oxiridagi qoldiq:"), None, bal])
        for r in (ws.max_row - 1, ws.max_row):
            for col in (3, 4):
                ws.cell(r, col).number_format = money
                ws.cell(r, col).font = bold
        widths = [12, 9, 16, 16, 38, 13, 24, 9, 70]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    wb.save(path)
    return {
        "file": Path(path).name, "layout": layout, "language": lang, "header_rows": header_rows,
        "first_data_row": first, "opening_balance": opening, "closing_balance": bal,
        "period": {"start": start.isoformat(), "end": end.isoformat()},
        "month_end_balances": month_end,
        "rows": [{"row": r, "date": t.d.isoformat(), "amount": t.amount, "category": t.category,
                  "recurring_group": t.recurring_group, "counterparty_stir": t.party.stir} for r, t in rows_truth],
    }


# ---------------------------------------------------------------- monthly simulation helpers

def split_amount(rng, total: int, n: int, min_part: int = 50_000, exact: bool = False) -> list[int]:
    """Split `total` into n positive parts rounded to 1 000 (exactly n when `exact`)."""
    if not exact:
        n = max(1, min(n, max(1, total // max(min_part, 1))))
    w = rng.pareto(1.6, n) + 0.3
    parts = np.maximum(np.floor(total * w / w.sum() / 1000) * 1000, 1000)
    parts[np.argmax(parts)] += total - parts.sum()
    return [int(p) for p in parts]


@dataclass
class StatementCtx:
    company: Company
    rng: np.random.Generator
    grammar: Grammar
    customers: list[Party]
    suppliers: list[Party]
    landlord: Party
    utilities: dict[str, Party]
    lenders: dict[str, Party]
    rent_contract: str
    own: Party
    tax: Party
    fund: Party
    bank: Party


def make_ctx(rng, company: Company, lang: str = "ru", n_customers: int | None = None,
             n_suppliers: int | None = None, customers: list[Party] | None = None) -> StatementCtx:
    nc = n_customers or int(rng.integers(4, 26))
    ns = n_suppliers or int(rng.integers(3, 18))
    cust = customers or [party_from(rng, *counterparty(rng)) for _ in range(nc)]
    sup = [party_from(rng, *counterparty(rng)) for _ in range(ns)]
    land = party_from(rng, *counterparty(rng, "construction"))
    util = {"electricity": party_from(rng, "Hududiy elektr tarmoqlari AJ", "200933985"),
            "gas": party_from(rng, "Hududgaz ta'minot", "201055108"),
            "water": party_from(rng, "Suvsoz MChJ", "203540111")}
    own = Party(company.name, company.stir, account(rng, "20208"), company.mfo)
    return StatementCtx(company, rng, Grammar(rng, lang), cust, sup, land, util, {},
                        f"{int(rng.integers(1, 99))}/{int(rng.integers(20, 26))}", own,
                        Party(TAX_OFFICE[0], TAX_OFFICE[1], "23402000300100001010", "00014"),
                        Party(FUND[0], FUND[1], "23414000400100001020", "00014"),
                        Party(company.bank, None, "45994000900000001001", company.mfo))


def month_txns(ctx: StatementCtx, y: int, m: int, *, sales: int, suppliers: int, payroll: int = 0,
               rent: int = 0, electricity: int = 0, gas: int = 0, water: int = 0, vat: int = 0,
               profit_tax: int = 0, loans: list[dict] | None = None, bank_fee: int = 0,
               cash_withdrawal: int = 0, vague: int = 0, other: int = 0, avg_receipt: int | None = None,
               avg_payment: int | None = None, salary_day: int = 5, customer_weights=None,
               supplier_weights=None, extra: list[Txn] | None = None, n_sales: int | None = None,
               n_suppliers: int | None = None, payroll_net: int | None = None,
               payroll_taxes: tuple[int, int] | None = None) -> list[Txn]:
    r, g = ctx.rng, ctx.grammar
    out: list[Txn] = []
    if sales > 0:
        n = n_sales or max(1, int(round(sales / (avg_receipt or max(sales / r.uniform(6, 40), 1)))))
        for amt in split_amount(r, sales, n, exact=bool(n_sales)):
            w = customer_weights
            idx = int(r.choice(len(ctx.customers), p=w)) if w is not None else int(min(r.pareto(1.2), len(ctx.customers) - 1))
            p = ctx.customers[idx]
            d = business_day(r, y, m)
            out.append(Txn(d, amt, p, g.render("sales_receipt", d=d, month=m, year=y), "sales_receipt"))
    if suppliers > 0:
        n = n_suppliers or max(1, int(round(suppliers / (avg_payment or max(suppliers / r.uniform(4, 30), 1)))))
        for amt in split_amount(r, suppliers, n, exact=bool(n_suppliers)):
            w = supplier_weights
            p = ctx.suppliers[int(r.choice(len(ctx.suppliers), p=w)) if w is not None else int(min(r.pareto(1.0), len(ctx.suppliers) - 1))]
            d = business_day(r, y, m)
            out.append(Txn(d, -amt, p, g.render("supplier_payment", d=d, month=m, year=y), "supplier_payment"))
    if payroll > 0 or payroll_net:
        net = payroll_net or int(round(payroll * 0.88, -3))
        pm = m - 1 or 12
        py = y if m > 1 else y - 1
        d = business_day(r, y, m, salary_day, spread=1)
        out.append(Txn(d, -net, ctx.bank, g.render("salary", month=pm, year=py), "salary", group="salary"))
        d2 = business_day(r, y, m, 15, spread=1)
        soc, pit = payroll_taxes or (int(round(payroll * 0.12, -3)), int(round(payroll * 0.12, -3)))
        out.append(Txn(d2, -soc, ctx.tax, g.render("social_tax", month=pm, year=py), "payroll_tax", group="payroll:social"))
        out.append(Txn(d2, -pit, ctx.tax, g.render("pit", month=pm, year=py), "payroll_tax", group="payroll:pit"))
    if rent > 0:
        d = business_day(r, y, m, 8, spread=1)
        out.append(Txn(d, -rent, ctx.landlord, g.render("rent", month=m, year=y, contract=ctx.rent_contract), "rent_utilities", group="rent"))
    for kind, amt in (("electricity", electricity), ("gas", gas), ("water", water)):
        if amt > 0:
            d = business_day(r, y, m, 25, spread=2)
            out.append(Txn(d, -amt, ctx.utilities[kind], g.render(kind, month=m, year=y), "rent_utilities", group=f"util:{kind}"))
    if vat > 0:
        d = business_day(r, y, m, 22, spread=1)
        out.append(Txn(d, -vat, ctx.tax, g.render("vat", month=(m - 1 or 12), year=y if m > 1 else y - 1), "tax", group="tax:vat"))
    if profit_tax > 0:
        d = business_day(r, y, m, 23, spread=1)
        out.append(Txn(d, -profit_tax, ctx.tax, g.render("profit_tax", month=m, year=y), "tax", group="tax:profit"))
    for ln in loans or []:
        p = ln["party"]
        d = ln.get("date") or business_day(r, y, m, ln["day"], spread=0)
        if ln.get("split"):
            out.append(Txn(d, -ln["principal"], p, g.render("loan_repayment", contract=ln["contract"], contract_date=ln.get("contract_date", "")),
                           "loan_repayment", group=f"loan:{ln['contract']}"))
            out.append(Txn(d, -ln["interest"], p, g.render("loan_interest", contract=ln["contract"]), "loan_repayment",
                           group=f"loan:{ln['contract']}:interest"))
        else:
            out.append(Txn(d, -ln["amount"], p, ln.get("purpose") or g.render("loan_repayment", contract=ln["contract"], contract_date=ln.get("contract_date", "")),
                           "loan_repayment", group=f"loan:{ln['contract']}"))
    if bank_fee > 0:
        d = business_day(r, y, m, calendar.monthrange(y, m)[1])
        out.append(Txn(d, -bank_fee, ctx.bank, g.render("bank_fee", month=m, year=y), "bank_fee"))
    if cash_withdrawal > 0:
        d = business_day(r, y, m)
        out.append(Txn(d, -cash_withdrawal, ctx.bank, g.render("cash_withdrawal"), "cash_withdrawal"))
    for _ in range(vague):
        d = business_day(r, y, m)
        p = ctx.suppliers[int(r.integers(0, len(ctx.suppliers)))]
        amt = int(round(r.uniform(0.5, 8.0) * 1e6, -3))
        cat = str(r.choice(["supplier_payment", "other"], p=[0.7, 0.3]))
        out.append(Txn(d, -amt, p, g.render("vague", d=d), cat, meta={"vague": True}))
    for _ in range(other):
        d = business_day(r, y, m)
        out.append(Txn(d, -int(round(r.uniform(0.3, 5.0) * 1e6, -3)), party_from(r, *counterparty(r)), g.render("other", d=d), "other"))
    out.extend(extra or [])
    return out


def balance_to_target(txns: list[Txn], opening: int, target: int, until: date, rng, since: date | None = None) -> None:
    """Scale supplier payments in (since, until] so that the balance at `until` equals `target`."""
    flows = sum(t.amount for t in txns if t.d <= until)
    diff = (opening + flows) - target           # >0: too much cash -> pay suppliers more
    sup = [t for t in txns if t.category == "supplier_payment" and t.d <= until and not t.meta.get("fixed")
           and (since is None or t.d > since)]
    total = -sum(t.amount for t in sup)
    if not sup or total <= 0:
        return
    k = (total + diff) / total
    k = max(k, 0.2)
    acc = 0
    for t in sup:
        t.amount = -int(round(-t.amount * k, -3))
        acc += t.amount
    # put the rounding residue on the largest payment
    flows2 = sum(t.amount for t in txns if t.d <= until)
    resid = (opening + flows2) - target
    big = min(sup, key=lambda t: t.amount)
    big.amount -= resid
