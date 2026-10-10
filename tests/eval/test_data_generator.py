"""Generator invariants: statutory sums hold, hero figures match the design, truth matches the files."""
import csv
import json
from datetime import date
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from eval.generator.company import make_company
from eval.generator.finmodel import make_fin
from eval.generator.heroes import FAR_F1, FAR_F2, _check_farovon, zilola_f1
from eval.generator.spec import check_sums
from eval.generator.statement import assign_recurring, make_ctx, month_txns, write_statement
from eval.generator.totals import f1_totals, f2_totals
from eval.generator.translit import to_cyrillic

ROOT = Path(__file__).resolve().parents[2]
SEED = ROOT / "data" / "seed"


@pytest.mark.parametrize("seed", range(12))
def test_financial_model_balances_and_sums(seed):
    rng = np.random.default_rng(seed)
    fin = make_fin(rng, make_company(rng))
    for d, v in fin.f1.items():
        assert v["400"] == v["780"], d
        assert check_sums("form1", v) == [], d
    for p, v in fin.f2.items():
        assert check_sums("form2", v) == [], p
    # retained earnings roll exactly by net profit (R04 holds on clean bundles)
    assert fin.f1["2026-06-30"]["450"] - fin.f1["2025-12-31"]["450"] == fin.f2["H1-2026"]["270"]
    assert fin.f1["2025-12-31"]["450"] - fin.f1["2024-12-31"]["450"] == fin.f2["FY2025"]["270"]


def test_farovon_matches_design():
    f1 = {k: f1_totals(v) for k, v in FAR_F1.items()}
    f2 = {k: f2_totals(v) for k, v in FAR_F2.items()}
    got = _check_farovon(f1, f2)          # asserts every pinned number internally
    assert got["dscr_new"] == 1.47 and got["dscr_incl_undeclared"] == 1.04
    assert f2["FY2025"]["010"] == 9_842_316 and f2["H1-2026"]["010"] == 5_316_428


def test_zilola_monitoring_numbers():
    f1, f2 = zilola_f1()
    q2, q3 = f2["Q2-2026"], f2["Q3-2026"]
    assert round((q3["010"] / q2["010"] - 1) * 100, 1) == -28.0
    assert q3["270"] == -102_000
    assert f1["2026-09-30"]["220"] / f1["2026-06-30"]["220"] - 1 >= 0.25
    assert f1["2026-09-30"]["210"] / f1["2026-06-30"]["210"] - 1 >= 0.25
    assert round(f1["2026-09-30"]["340"] / 46_800, 1) == 1.4
    for v in f1.values():
        assert v["400"] == v["780"]


def test_statement_truth_matches_xlsx(tmp_path):
    rng = np.random.default_rng(5)
    co = make_company(rng)
    ctx = make_ctx(rng, co)
    txns = []
    for m in (1, 2, 3):
        txns += month_txns(ctx, 2026, m, sales=500_000_000, suppliers=420_000_000, payroll=40_000_000, rent=5_000_000,
                           bank_fee=300_000, n_sales=20, n_suppliers=10)
    assign_recurring(txns)
    for layout in ("A", "B"):
        st = write_statement(tmp_path / f"s{layout}.xlsx", co, txns, 100_000_000, date(2026, 1, 1), date(2026, 3, 31), layout=layout)
        ws = load_workbook(tmp_path / f"s{layout}.xlsx", data_only=True).active
        for r in st["rows"]:
            debit, credit = ws.cell(r["row"], 3).value, ws.cell(r["row"], 4).value
            assert (credit or 0) - (debit or 0) == r["amount"]
        assert sum(1 for r in st["rows"] if r["recurring_group"] == "salary") == 3


def test_translit():
    assert to_cyrillic("II bo'lim bo'yicha jami") == "II бўлим бўйича жами"
    assert to_cyrillic("Farovon Savdo MChJ") == "Фаровон Савдо МЧЖ"


@pytest.mark.skipif(not (SEED / "loan_book.csv").exists(), reason="seed data not generated")
def test_seed_loan_book():
    rows = list(csv.DictReader(open(SEED / "loan_book.csv", encoding="utf-8")))
    assert len(rows) == 148
    assert len({r["loan_id"] for r in rows}) == 148
    out = sorted(float(r["outstanding"]) for r in rows)
    assert round(sum(out) / 1000, 1) == 86.4
    assert float(np.median(out)) == 412.0 and max(out) == 4200.0
    for r in rows:
        assert (SEED / "portfolio" / r["loan_id"] / "truth.json").exists()


@pytest.mark.skipif(not (SEED / "hero_farovon" / "truth.json").exists(), reason="seed data not generated")
def test_hero_truth_files():
    t = json.loads((SEED / "hero_farovon" / "truth.json").read_text(encoding="utf-8"))
    assert t["expected_findings_exactly"] == ["R06", "R07", "R10"]
    assert t["statement_summary"]["rows"] == 1946
    for d in t["documents"]:
        assert (SEED / "hero_farovon" / d["file"]).exists()
    names = {d["file"] for d in t["documents"]}
    assert {"Ariza_FarovonSavdo.pdf", "Balans_2025.pdf", "Balans_1H2026.pdf", "Forma2_2025.pdf", "Forma2_1H2026.pdf",
            "vypiska_2026.xlsx"} == names
