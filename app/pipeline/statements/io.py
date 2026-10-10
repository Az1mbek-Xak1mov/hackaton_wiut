"""Minimal xlsx loader for tests and the categoriser training script (the pipeline uses parse_xlsx)."""
from __future__ import annotations

from pathlib import Path

from app.pipeline.types import ParsedDoc, XlsxSheet


def load_xlsx_sheets(path: str | Path) -> list[XlsxSheet]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        return [XlsxSheet(name=ws.title, rows=[list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()


def load_xlsx_doc(path: str | Path) -> ParsedDoc:
    return ParsedDoc(kind="xlsx", doc_type="bank_statement", sheets=load_xlsx_sheets(path))
